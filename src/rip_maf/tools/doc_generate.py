"""doc.generate tool — report files AND verbatim text files.

Two paths, exactly one per call:

1. Report path (``title`` + ``sections[{heading, body}]``, optional
   ``tables``): renders Markdown + DOCX + PDF from one template model.
   Markdown renders with the stdlib; DOCX (`python-docx`) and PDF
   (`reportlab`) are LAZY imports — a deployment without them fails honest
   per call instead of breaking tool import.
2. Verbatim path (``content`` + ``format``/``filename``): writes the text
   byte-for-byte as a txt/csv/md/json/code file (no LLM rendering, no
   report template). The writer's prose is the file — e.g. a reasoning
   step's CSV table becomes ``result.csv``.

Binaries travel as base64 in ToolResponse.data with Markdown/text inline
as `output`; durable delivery arrives with file-based artifacts (Q34).

Effect class: sandboxed (bounded compute producing artifacts).

RIP port: no plugin system — direct Tool subclass (ADR-017).
"""

from __future__ import annotations

import base64
import io
from typing import Any, ClassVar

from rip_maf.tools.base import Tool, ToolRequest, ToolResponse


def _need_modules() -> tuple[Any, Any]:
    """Lazily import the DOCX/PDF toolchains (ImportError stays catchable)."""
    try:
        from docx import Document  # type: ignore[import-not-found]
    except ImportError as e:
        raise RuntimeError(f"python-docx is not installed: {e}") from e
    try:
        from reportlab.lib.pagesizes import letter  # type: ignore[import-not-found]
        from reportlab.lib.styles import (
            getSampleStyleSheet,  # type: ignore[import-not-found]
        )
        from reportlab.platypus import (  # type: ignore[import-not-found]
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as e:
        raise RuntimeError(f"reportlab is not installed: {e}") from e
    return Document, (
        letter,
        getSampleStyleSheet,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )


def render_markdown(title: str, sections: list[dict], tables: list[dict]) -> str:
    lines = [f"# {title}", ""]
    for section in sections:
        lines += [f"## {section['heading']}", "", section["body"], ""]
    for table in tables:
        lines.append(" | ".join(table["headers"]))
        lines.append(" | ".join("---" for _ in table["headers"]))
        for row in table["rows"]:
            lines.append(" | ".join(str(cell) for cell in row))
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_docx(title: str, sections: list[dict], tables: list[dict]) -> bytes:
    Document, _ = _need_modules()
    doc = Document()
    doc.add_heading(title, level=0)
    for section in sections:
        doc.add_heading(section["heading"], level=1)
        doc.add_paragraph(section["body"])
    for table in tables:
        grid = [table["headers"], *table["rows"]]
        doc_table = doc.add_table(rows=len(grid), cols=len(table["headers"]))
        for i, row in enumerate(grid):
            for j, cell in enumerate(row):
                doc_table.cell(i, j).text = str(cell)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(title: str, sections: list[dict], tables: list[dict]) -> bytes:
    from html import escape

    from reportlab.lib import colors  # type: ignore[import-not-found]

    _, mods = _need_modules()
    (
        letter,
        getSampleStyleSheet,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    ) = mods
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    for section in sections:
        story += [
            Paragraph(escape(section["heading"]), styles["Heading2"]),
            Paragraph(escape(section["body"]).replace("\n", "<br/>"), styles["Normal"]),
            Spacer(1, 6),
        ]
    for table in tables:
        grid = [
            [escape(str(c)) for c in row] for row in [table["headers"], *table["rows"]]
        ]
        wrapped = [[Paragraph(c, styles["Normal"]) for c in row] for row in grid]
        story.append(
            Table(
                wrapped,
                style=TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ]
                ),
            )
        )
        story.append(Spacer(1, 6))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter).build(story)
    return buf.getvalue()


def _valid_sections(sections: Any) -> list[dict] | None:
    if not isinstance(sections, list) or not sections:
        return None
    clean: list[dict] = []
    for item in sections:
        if (
            not isinstance(item, dict)
            or not isinstance(item.get("heading"), str)
            or not isinstance(item.get("body"), str)
        ):
            return None
        clean.append({"heading": item["heading"], "body": item["body"]})
    return clean


def _valid_tables(tables: Any) -> list[dict] | None:
    if tables is None:
        return []
    if not isinstance(tables, list):
        return None
    clean: list[dict] = []
    for item in tables:
        if (
            not isinstance(item, dict)
            or not isinstance(item.get("headers"), list)
            or not item["headers"]
            or not isinstance(item.get("rows"), list)
        ):
            return None
        clean.append(
            {"headers": [str(h) for h in item["headers"]], "rows": item["rows"]}
        )
    return clean


#: Plain-text extensions the verbatim path may write. Deliberately
#: data/text only — no executables, no archives, no rendered binaries
#: (those stay on the report path or doc.convert).
_VERBATIM_EXTENSIONS = frozenset({
    "txt", "md", "csv", "json", "py", "js", "ts", "html", "css", "sh",
    "yaml", "yml", "xml",
})

_VERBATIM_MIMES = {
    "txt": "text/plain",
    "md": "text/markdown",
    "csv": "text/csv",
    "json": "application/json",
    "py": "text/x-python",
    "js": "text/javascript",
    "ts": "text/typescript",
    "html": "text/html",
    "css": "text/css",
    "sh": "text/x-sh",
    "yaml": "text/yaml",
    "yml": "text/yaml",
    "xml": "text/xml",
}


def _verbatim_extension(tool_input: dict) -> tuple[str | None, str | None]:
    """Resolve (extension, error) for the verbatim path.

    `filename` (e.g. "result.csv") wins when present; else `format`
    (default "txt"). Both present must agree — a caller saying
    format=csv for "notes.txt" is confused, fail honest.
    """
    raw_name = tool_input.get("filename")
    name_ext: str | None = None
    if raw_name is not None:
        name = str(raw_name).strip()
        if "/" in name or "\\" in name or not name or name.startswith("."):
            return None, "'filename' must be a bare file name like 'result.csv'"
        ext = name.rsplit(".", 1)[1].lower() if "." in name else ""
        if ext not in _VERBATIM_EXTENSIONS:
            return None, (
                f"'filename' extension must be one of {sorted(_VERBATIM_EXTENSIONS)} "
                f"(got {ext!r})"
            )
        name_ext = ext
    raw_format = tool_input.get("format")
    fmt: str | None = None
    if raw_format is not None:
        fmt = str(raw_format).strip().lower()
        if fmt not in _VERBATIM_EXTENSIONS:
            return None, (
                f"'format' must be one of {sorted(_VERBATIM_EXTENSIONS)} "
                f"(got {fmt!r})"
            )
    if name_ext is not None and fmt is not None and name_ext != fmt:
        return None, (
            f"'filename' (.{name_ext}) and 'format' ({fmt!r}) disagree — "
            "pass one, or make them agree"
        )
    return name_ext or fmt or "txt", None


class DocGenerateTool(Tool):
    tool_id = "doc.generate"
    name = "Doc Generate"
    description = (
        "Write a file: either a titled report (title + sections[{heading, body}], "
        "optional tables → Markdown/DOCX/PDF) or a verbatim text file "
        "(content + format/filename → txt/csv/md/json/code, byte-for-byte). "
        'Report example: {"title":"Docker Overview","sections":[{"heading":"Intro","body":"..."}]}. '
        'Verbatim example: {"content":"a,b\\n1,2\\n","filename":"result.csv"}. '
        "Exactly one of sections/content per call."
    )
    input_schema: ClassVar[dict] = {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "sections": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {"type": "string"},
                        "body": {"type": "string"},
                    },
                    "required": ["heading", "body"],
                },
            },
            "tables": {"type": "array"},
            "content": {"type": "string"},
            "format": {"type": "string"},
            "filename": {"type": "string"},
        },
        "required": [],
        "additionalProperties": False,
    }
    input_example: ClassVar[str] = (
        'doc.generate {"title": "Docker Overview", "sections": '
        '[{"heading": "Intro", "body": "..."}], "tables": []}. A section '
        "body may be a {{id}} placeholder carrying an upstream answer step's "
        "prose. Verbatim file instead: "
        'doc.generate {"content": "a,b\\n1,2\\n", "filename": "result.csv"}.'
    )
    output_schema: ClassVar[dict] = {
        "type": "object",
        "properties": {
            "markdown": {"type": "string"},
            "docx_b64": {"type": "string"},
            "pdf_b64": {"type": "string"},
        },
    }
    effect_class = "sandboxed"  # type: ignore[assignment]
    cost_class = "medium"

    def validate_input(self, tool_input: dict) -> str | None:
        # Path presence first (before the closed-key schema check) so a
        # call on neither path names both paths instead of one stray key.
        has_sections = (
            "sections" in tool_input and tool_input["sections"] is not None
        )
        has_content = (
            "content" in tool_input and tool_input["content"] is not None
        )
        if has_sections and has_content:
            return (
                "pass either 'sections' (titled report) or 'content' "
                "(verbatim file), never both"
            )
        if not has_sections and not has_content:
            return (
                "pass either 'sections' (titled report: title + sections) or "
                "'content' (verbatim file: content + format/filename)"
            )
        err = super().validate_input(tool_input)
        if err is not None:
            return err
        if has_sections:
            for stray in ("content", "format", "filename"):
                if stray in tool_input and tool_input[stray] is not None:
                    return (
                        f"'{stray}' belongs to the verbatim file path — "
                        "omit it for a titled report"
                    )
            if (
                not isinstance(tool_input.get("title"), str)
                or not tool_input["title"].strip()
            ):
                return "'title' is required in input"
            if _valid_sections(tool_input.get("sections")) is None:
                return "'sections' must be a non-empty array of {heading, body}"
            if _valid_tables(tool_input.get("tables")) is None:
                return "'tables' must be an array of {headers[], rows[][]}"
            return None
        if has_content:
            for stray in ("sections", "tables"):
                if stray in tool_input and tool_input[stray] is not None:
                    return (
                        f"'{stray}' belongs to the titled report path — "
                        "omit it for a verbatim file"
                    )
            if (
                not isinstance(tool_input.get("content"), str)
                or not tool_input["content"]
            ):
                return "'content' must be a non-empty string"
            _ext, ext_err = _verbatim_extension(tool_input)
            if ext_err is not None:
                return ext_err
            return None
        return (
            "pass either 'sections' (titled report: title + sections) or "
            "'content' (verbatim file: content + format/filename)"
        )

    def execute(self, request: ToolRequest) -> ToolResponse:
        invalid = self.invalid_response(request.input)
        if invalid.error is not None:
            return invalid
        if (
            "content" in request.input
            and request.input["content"] is not None
        ):
            return self._execute_verbatim(request.input)
        title = str(request.input["title"])
        sections = _valid_sections(request.input["sections"])
        tables = _valid_tables(request.input.get("tables"))
        markdown = render_markdown(title, sections, tables)
        try:
            docx_bytes = render_docx(title, sections, tables)
            pdf_bytes = render_pdf(title, sections, tables)
        except RuntimeError as e:
            return ToolResponse(
                tool_id=self.tool_id, ok=False, output=None, error=str(e)
            )
        except Exception as e:  # noqa: BLE001 - render failure is a tool failure
            return ToolResponse(
                tool_id=self.tool_id, ok=False, output=None, error=f"render failed: {e}"
            )
        return ToolResponse(
            tool_id=self.tool_id,
            ok=True,
            output=markdown,
            data={
                "markdown": markdown,
                "docx_b64": base64.b64encode(docx_bytes).decode("ascii"),
                "pdf_b64": base64.b64encode(pdf_bytes).decode("ascii"),
            },
        )

    def _execute_verbatim(self, tool_input: dict) -> ToolResponse:
        """Write `content` byte-for-byte; no template, no rendering."""
        content = str(tool_input["content"])
        ext, ext_err = _verbatim_extension(tool_input)
        if ext_err is not None:  # validated above; fail honest anyway
            return ToolResponse(
                tool_id=self.tool_id, ok=False, output=None, error=ext_err
            )
        raw_name = tool_input.get("filename")
        filename = (
            str(raw_name).strip()
            if isinstance(raw_name, str) and str(raw_name).strip()
            else f"document.{ext}"
        )
        mime = _VERBATIM_MIMES[ext]
        return ToolResponse(
            tool_id=self.tool_id,
            ok=True,
            output=content,
            data={
                "file_b64": base64.b64encode(content.encode("utf-8")).decode("ascii"),
                "filename": filename,
                "mime": mime,
            },
        )
